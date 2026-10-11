"""Mise en page et génération du .pptx au design système (Poppins, blanc, #122F4C, fond dégradé)."""
from __future__ import annotations

import copy
import os
import re

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
TEMPLATE = os.path.join(ASSETS, "template.pptx")

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x12, 0x2F, 0x4C)

# Une seule famille « Poppins » (Regular / Italic / Bold) : c'est le nom que PowerPoint, Keynote
# et Canva reconnaissent partout. Les graisses Light/Medium sont souvent mal reprises à l'import.
FONT_FILES = {"L": "Poppins-Regular.ttf", "LI": "Poppins-Italic.ttf", "M": "Poppins-Bold.ttf", "B": "Poppins-Bold.ttf"}
FONTS = {k: ImageFont.truetype(os.path.join(ASSETS, "fonts", v), 100) for k, v in FONT_FILES.items()}

FACE = {"L": ("Poppins", False, False), "LI": ("Poppins", False, True),
        "M": ("Poppins", True, False), "B": ("Poppins", True, False)}

# styles : taille (pt), graisse, interligne (x taille), espacement des lettres (pt), majuscules
STY = {
    "eyebrow":  dict(size=15, st="B", lh=1.35, cs=1.4, caps=True),
    "ref":      dict(size=16, st="B", lh=1.35, cs=1.0, caps=True),
    "section":  dict(size=18, st="B", lh=1.35, cs=1.2, caps=True),
    "cover":    dict(size=46, st="M", lh=1.2),
    "hero":     dict(size=38, st="L", lh=1.25),
    "stmt":     dict(size=28, st="L", lh=1.35),
    "stmt_l":   dict(size=32, st="L", lh=1.3),
    "intro":    dict(size=20, st="LI", lh=1.4),
    "big":      dict(size=60, st="L", lh=1.15),
    "quote":    dict(size=24, st="L", lh=1.45),
    "quote_l":  dict(size=28, st="L", lh=1.4),
    "quote_xl": dict(size=36, st="L", lh=1.3),
    "label":    dict(size=16, st="LI", lh=1.3),
    "q":        dict(size=28, st="B", lh=1.3),
    "a":        dict(size=28, st="L", lh=1.3),
}

LAYOUT = {"cover": "DECLARATION", "statement": "DECLARATION", "reading": "LECTURE", "plan": "LECTURE",
          "question": "INTERPELLATION", "twocol": "DEUX_COLONNES"}


# ------------------------------------------------------------------ typographie
def nb(t: str) -> str:
    """Espaces insécables de la typographie française."""
    t = re.sub(r"«\s+", "« ", t)
    t = re.sub(r"\s+»", " »", t)
    return re.sub(r" ([;:?!])", " \\1", t)


def tw(text, st, size, cs=0.0):
    return FONTS[st].getlength(text) * size / 100 + cs * len(text)


def count_lines(runs, size, width, cs=0.0):
    lines, cur = 1, 0.0
    for text, st in runs:
        parts = text.split(" ")
        for i, word in enumerate(parts):
            w = tw(word, st, size, cs)
            sp = tw(" ", st, size, cs) if i < len(parts) - 1 else 0
            if cur > 0 and cur + w > width:
                lines += 1
                cur = 0.0
            cur += w + sp
    return lines


def split_emphasis(text: str, emphasis: list[str]) -> list[tuple[str, bool]]:
    """Découpe `text` en segments (texte, en_valeur)."""
    marks = [False] * len(text)
    for e in emphasis or []:
        e = e.strip()
        if not e:
            continue
        start = text.find(e)
        if start < 0:
            # tolère les espaces insécables / différences d'espaces
            pat = re.escape(e).replace(r"\ ", r"[\s ]+")
            m = re.search(pat, text)
            if not m:
                continue
            start, end = m.span()
        else:
            end = start + len(e)
        for i in range(start, end):
            marks[i] = True
    segs = []
    for ch, m in zip(text, marks):
        if segs and segs[-1][1] == m:
            segs[-1] = (segs[-1][0] + ch, m)
        else:
            segs.append((ch, m))
    return segs or [("", False)]


# ------------------------------------------------------------------ choix des styles
def style_for(mode: str, btype: str, slide: dict) -> str | None:
    quotes = sum(len(b["text"]) for b in slide.get("blocks", []) if b["type"] == "quote")
    stmts = sum(len(b["text"]) for b in slide.get("blocks", []) if b["type"] in ("statement", "title"))
    if btype == "eyebrow":
        return "eyebrow"
    if btype == "ref":
        return "ref"
    if btype == "section":
        return "section"
    if btype == "intro":
        return "intro"
    if btype == "big":
        return "big"
    if btype == "label":
        return "label"
    if mode == "cover":
        return {"title": "cover", "statement": "hero"}.get(btype, "hero")
    if mode == "question":
        return "stmt_l" if stmts < 260 else "stmt"
    if mode == "reading":
        if btype == "question":
            return "q"
        if btype == "answer":
            return "a"
        return "quote_l" if quotes < 300 else "quote"
    # statement
    if btype == "title":
        return "hero"
    if btype == "statement":
        return "stmt_l" if stmts < 220 else "stmt"
    if btype == "question":
        return "q"
    if btype == "answer":
        return "a"
    if btype == "quote":
        return "quote_xl" if quotes < 120 else "quote_l"
    return "stmt"


def gap(prev: str | None, cur: str) -> float:
    if prev is None:
        return 0
    table = {
        ("dash", None): 22, ("section", "dash"): 20, ("eyebrow", "dash"): 20,
        ("section", "ref"): 20, ("section", None): 20,
        ("ref", "quote"): 6, ("ref", "label"): 14, ("label", "quote"): 4,
        ("quote", "ref"): 26, ("quote", "quote"): 12, ("quote", "label"): 26,
        ("question", "answer"): 4, ("intro", None): 34, ("big", "big"): 12,
        ("statement", "statement"): 14, ("title", "title"): 40, ("title", "statement"): 22,
        ("statement", "title"): 22, ("eyebrow", "title"): 22, ("title", "footer"): 0,
    }
    if cur == "question":
        return 40
    return table.get((prev, cur), table.get((prev, None), 18))


# ------------------------------------------------------------------ écriture pptx
def _textbox(slide, x, y, w, h, runs, size, lh, cs, color, name, alpha=None):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.line_spacing = Pt(lh)
    for text, st, *rest in runs:
        r = p.add_run()
        r.text = text
        face, bold, italic = FACE[st]
        f = r.font
        f.name, f.size, f.bold, f.italic = face, Pt(size), bold, italic
        f.color.rgb = color
        rpr = r._r.get_or_add_rPr()
        if cs:
            rpr.set("spc", str(int(cs * 100)))
        run_alpha = rest[0] if rest else alpha
        if run_alpha is not None:
            clr = rpr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = clr.makeelement(qn("a:alpha"), {"val": str(int(run_alpha * 1000))})
            clr.append(a)
        # polices latines + secours
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", face)
    return tb


def _dash(slide, x, y, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(40), Pt(3))
    sh.name = "Trait"
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False


def _flow(slide_spec):
    """Calcule la position de chaque bloc ; réduit tout si ça ne tient pas."""
    mode = slide_spec["mode"]
    in_card = mode in ("reading", "question")
    x, w = (80, 790) if in_card else (64, 800)
    top, bottom = (64, 476) if in_card else (56, 484)
    blocks = [b for b in slide_spec.get("blocks", []) if b.get("type") != "footer"]
    scale = 1.0
    while True:
        items, cur, prev = [], 0.0, None
        for b in blocks:
            t = b["type"]
            cur += gap(prev, t) * scale
            prev = t
            if t == "dash":
                items.append(dict(kind="dash", y=cur)); cur += 3; continue
            sname = style_for(mode, t, slide_spec)
            st = STY[sname]
            size = round(st["size"] * scale, 1)
            text = nb(b.get("text", "").strip())
            if st.get("caps"):
                text = text.upper()
                emph = [e.upper() for e in b.get("emphasis") or []]
            else:
                emph = b.get("emphasis") or []
            base = st["st"]
            if t == "question" and mode == "question":
                runs = [(text, "B")]
            elif sname == "q":
                runs = [(text, "B")]
            else:
                runs = [(seg, "B" if m else base) for seg, m in split_emphasis(text, [nb(e) for e in emph])]
            n = count_lines(runs, size, w - 6, st.get("cs", 0))
            lh = size * st["lh"]
            items.append(dict(kind="text", y=cur, h=n * lh + 2, runs=runs, size=size, lh=lh, cs=st.get("cs", 0), name=t))
            cur += n * lh
        total = cur
        if total <= bottom - top or scale <= 0.62:
            break
        scale -= 0.04
    if mode == "cover":
        start = 75
    elif in_card:
        start = max(top, 270 - total / 2)
    else:
        start = max(top, 255 - total / 2)
    start = min(start, max(top, bottom - total))
    for it in items:
        it["y"] += start
        it["x"], it["w"] = x, w
    return items


def build(slides: list[dict], out_path: str) -> str:
    prs = Presentation(TEMPLATE)
    layouts = {l.name: l for l in prs.slide_layouts}
    for spec in slides:
        mode = spec.get("mode", "statement")
        if mode not in LAYOUT:
            mode = "statement"
        spec["mode"] = mode
        slide = prs.slides.add_slide(layouts[LAYOUT[mode]])
        ink = WHITE

        if mode == "plan":
            rows = spec.get("rows") or []
            _textbox(slide, 80, 60, 790, 56, [("Plan", "B")], 40, 50, 0, WHITE, "Titre")
            # taille commune à toutes les lignes, réduite si le plan est long
            def plan_runs(r):
                runs = [(nb(r.get("text", "")), "L")]
                if r.get("ref"):
                    runs.append(("  ·  " + nb(r["ref"]), "B", 75))  # références toujours en gras
                return runs
            size = 24
            while True:
                lines = [count_lines([(t, st) for t, st, *_ in plan_runs(r)], size, 730) for r in rows]
                total = sum(n * size * 1.3 for n in lines) + 12 * max(0, len(rows) - 1)
                if total <= 350 or size <= 16:
                    break
                size -= 1
            y = 136
            for i, (r, n) in enumerate(zip(rows, lines), 1):
                h = n * size * 1.3
                _textbox(slide, 80, y, 54, h, [(r.get("num", ""), "B")], size * 0.85, size * 1.3, 0, WHITE, f"Num {i}")
                _textbox(slide, 140, y, 740, h, plan_runs(r), size, size * 1.3, 0, WHITE, f"Partie {i}")
                y += h + 12
            continue

        if mode == "twocol":
            for k, c in enumerate((spec.get("cols") or [])[:2]):
                cx = 40 + k * 450
                body = nb(c.get("body", ""))
                size = 20
                while count_lines([(body, "L")], size, 366 - 6) * size * 1.5 > 320 and size > 13:
                    size -= 0.5
                _textbox(slide, cx + 32, 90, 366, 42, [(nb(c.get("head", "")), "B")], 30, 38, 0, WHITE, f"Titre colonne {k + 1}")
                _textbox(slide, cx + 32, 144, 366, 320, [(body, "L")], size, size * 1.5, 0, WHITE, f"Texte colonne {k + 1}")
            continue

        for it in _flow(spec):
            if it["kind"] == "dash":
                _dash(slide, it["x"], it["y"], ink)
            else:
                _textbox(slide, it["x"], it["y"], it["w"], it["h"], it["runs"], it["size"], it["lh"], it["cs"], ink, it["name"])
        footer = [b for b in spec.get("blocks", []) if b.get("type") == "footer"]
        if footer:
            _textbox(slide, 64, 450, 760, 26, [(nb(footer[0]["text"]), "L")], 16, 22, 0, ink, "Pied de page")
    prs.save(out_path)
    return out_path
