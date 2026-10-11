"""Extraction du contenu d'un .pptx ou d'un .pdf, slide par slide.

Chaque slide devient un dict :
  {"index": 1, "background": "dark|light|yellow|image|unknown",
   "paragraphs": [{"text": "...", "size": 24.0, "bold": False, "emph": [...], "top": 120.0}]}
`emph` liste les morceaux du paragraphe mis en valeur (gras ou couleur différente).
"""
from __future__ import annotations

import pymupdf
from pptx import Presentation
from pptx.util import Emu


# --------------------------------------------------------------------------- PPTX
def _shape_bg_kind(slide) -> str:
    try:
        fill = slide.background.fill
        if fill.type == 1:  # solid
            return _classify_rgb(fill.fore_color.rgb)
        if fill.type == 6:
            return "image"
    except Exception:
        pass
    return "unknown"


def _classify_rgb(rgb) -> str:
    r, g, b = rgb[0], rgb[1], rgb[2]
    lum = 0.299 * r + 0.587 * g + 0.114 * b
    if lum < 90:
        return "dark"
    if r > 200 and g > 180 and b < 170:
        return "yellow"
    return "light"


def _run_color(run):
    try:
        if run.font.color and run.font.color.type is not None and run.font.color.rgb is not None:
            return str(run.font.color.rgb)
    except Exception:
        pass
    return None


def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # groupe
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def extract_pptx(path: str) -> list[dict]:
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        paras = []
        shapes = [s for s in _iter_shapes(slide.shapes) if s.has_text_frame]
        shapes.sort(key=lambda s: ((s.top or 0), (s.left or 0)))
        for sh in shapes:
            for p in sh.text_frame.paragraphs:
                runs = [r for r in p.runs if r.text]
                text = "".join(r.text for r in runs).strip()
                if not text:
                    continue
                sizes = [r.font.size.pt for r in runs if r.font.size]
                colors = [_run_color(r) for r in runs]
                main_color = max(set(colors), key=colors.count) if colors else None
                emph = [r.text for r, c in zip(runs, colors)
                        if (r.font.bold and not all(x.font.bold for x in runs)) or (c and c != main_color)]
                paras.append({
                    "text": text,
                    "size": round(max(sizes), 1) if sizes else None,
                    "bold": all(bool(r.font.bold) for r in runs),
                    "emph": [e.strip() for e in emph if e.strip()],
                    "top": round(Emu(sh.top or 0).pt),
                    "left": round(Emu(sh.left or 0).pt),
                })
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"index": i, "background": _shape_bg_kind(slide), "paragraphs": paras, "notes": notes})
    return slides


# --------------------------------------------------------------------------- PDF
def _page_bg(page) -> str:
    pix = page.get_pixmap(matrix=pymupdf.Matrix(0.1, 0.1))
    w, h = pix.width, pix.height
    samples = [pix.pixel(x, y) for x, y in ((1, 1), (w - 2, 1), (1, h - 2), (w - 2, h - 2), (w // 2, h - 2))]
    r = sum(s[0] for s in samples) // len(samples)
    g = sum(s[1] for s in samples) // len(samples)
    b = sum(s[2] for s in samples) // len(samples)
    return _classify_rgb((r, g, b))


def extract_pdf(path: str) -> list[dict]:
    doc = pymupdf.open(path)
    slides = []
    for i, page in enumerate(doc, 1):
        d = page.get_text("dict")
        lines = []
        for block in d["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if spans:
                    lines.append((line["bbox"], spans))
        # regroupe les lignes proches de même taille en paragraphes
        paras, cur = [], None
        for bbox, spans in sorted(lines, key=lambda l: (round(l[0][1]), l[0][0])):
            size = max(s["size"] for s in spans)
            text = "".join(s["text"] for s in spans)
            colors = [s["color"] for s in spans]
            if (cur and abs(cur["size"] - size) < 0.6 and abs(bbox[0] - cur["left"]) < 4
                    and bbox[1] - cur["_bottom"] < size * 0.9):
                cur["text"] = (cur["text"] + " " + text).strip()
                cur["lines"].append(text.strip())
                cur["_spans"] += spans
                cur["_bottom"] = bbox[3]
            else:
                cur = {"text": text.strip(), "size": round(size, 1), "top": round(bbox[1]), "left": round(bbox[0]),
                       "_bottom": bbox[3], "_spans": list(spans), "lines": [text.strip()]}
                paras.append(cur)
        for p in paras:
            spans = p.pop("_spans"); p.pop("_bottom")
            if len(p["lines"]) == 1:
                p.pop("lines")
            colors = [s["color"] for s in spans]
            main = max(set(colors), key=colors.count)
            bolds = ["bold" in s["font"].lower() or (s["flags"] & 16) for s in spans]
            p["bold"] = all(bolds)
            p["emph"] = [s["text"].strip() for s, b in zip(spans, bolds)
                         if s["text"].strip() and ((b and not p["bold"]) or s["color"] != main)]
        slides.append({"index": i, "background": _page_bg(page), "paragraphs": paras, "notes": ""})
    return slides


def extract(path: str) -> list[dict]:
    low = path.lower()
    if low.endswith(".pptx"):
        return extract_pptx(path)
    if low.endswith(".pdf"):
        return extract_pdf(path)
    raise ValueError("Format non pris en charge (attendu : .pptx ou .pdf)")
